"""Generate detailed Russian workbook, standalone HTML, Markdown and CSV; no macros."""
import csv
import json
import math
from pathlib import Path
import sys

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / '.tools/table-python'))
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.formatting.rule import FormulaRule
from network_test_guide import BASICS, REPORTS, SOLO, PREPARATION, ADDRESSES, TRANSPORTS, STEPS, DATA, FAILURES, SCENARIO_SETUP, NEGATIVE, EXTRA_STEPS
from network_test_render import render_guides

OUT = ROOT / 'docs/testing'
NAVY, INK = '17324D', '213547'

def csv_file(name, headers, rows):
    with (OUT / name).open('w', encoding='utf-8-sig', newline='') as file:
        writer = csv.writer(file, delimiter=';'); writer.writerow(headers); writer.writerows(rows)

def page(book, name, headers, rows, widths, table_name):
    sheet = book.create_sheet(name); sheet.append(headers)
    for row in rows: sheet.append(row)
    sheet.freeze_panes = 'C2' if name in ('Матрица','Прогоны','Сценарии подробно') else 'A2'
    sheet.sheet_view.showGridLines = False
    for col, width in enumerate(widths, 1): sheet.column_dimensions[sheet.cell(1,col).column_letter].width = width
    for cell in sheet[1]:
        cell.fill = PatternFill('solid', fgColor=NAVY)
        cell.font = Font(name='Calibri', bold=True, color='FFFFFF', size=11)
        cell.alignment = Alignment(wrap_text=True, vertical='center')
    sheet.row_dimensions[1].height = 36
    for row in sheet.iter_rows(min_row=2):
        lines = 1
        for index, cell in enumerate(row):
            cell.font = Font(name='Calibri', size=11, color=INK)
            cell.alignment = Alignment(wrap_text=True, vertical='top')
            lines = max(lines, sum(max(1, math.ceil(len(part) / max(4, widths[index] - 4))) for part in str(cell.value or '').split('\n')))
        sheet.row_dimensions[row[0].row].height = min(409, max(42, lines * 15 + 14))
    table = Table(displayName=table_name, ref=sheet.dimensions)
    table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2',showRowStripes=True)
    sheet.add_table(table)
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = 'landscape'; sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1; sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = '1:1'; sheet.oddFooter.center.text = 'SkyCraft YS | &P / &N'
    return sheet

def scenario_rows(scenarios):
    result = []
    for s in scenarios:
        sid = s['id']
        if sid in NEGATIVE:
            for index, (title, host, client, expected) in enumerate(NEGATIVE[sid],1):
                result.append([sid,str(index),title,host,client,expected])
            if sid in ('E01','E02'):
                result.append([sid,'Итог','Отчёт','Завершите лог и соберите ZIP по R2–R5. Роль Host только при debug start … host.', 'ID и роль как при debug start, затем восстановление G8. Для второго транспорта отдельный ID и отчёт.', s['expected']+' Требуемые минуты: '+str(s['minutes'])+'.'])
            continue
        condition, action = SCENARIO_SETUP[sid]
        result.append([sid,'0','Условия до прогона',condition,action,'Проверьте необходимые условия. При отсутствии доступного адреса/нужного прокси — BLOCKED с конкретной причиной.'])
        for number,title,host,client,expected in STEPS:
            if number == 'X1': continue
            host = host.replace('RUN_ID',sid+'_01').replace('I01_01',sid+'_01')
            client = client.replace('RUN_ID',sid+'_01').replace('I01_01',sid+'_01')
            if number == '8':
                host += f' Для {sid}: минимум {s["minutes"]} минут.'
                client += f' Для {sid}: минимум {s["minutes"]} минут.'
            result.append([sid,number,title,host,client,expected])
            for extra in EXTRA_STEPS.get(sid, {}).get(number, []):
                result.append([sid]+extra)
        result.append([sid,'Итог','Критерий сценария',s['expected'],s['limit'],'Неизвестные шаги = PARTIAL. Точное нарушение ожидания = FAIL.'])
    return result

def main():
    OUT.mkdir(parents=True,exist_ok=True)
    scenarios=json.loads((OUT/'scenarios.json').read_text(encoding='utf-8'))
    assert set(s['id'] for s in scenarios)==set(SCENARIO_SETUP)|set(NEGATIVE), 'Missing scenario instructions'
    for scenario in scenarios:
        sid = scenario['id']
        if sid in NEGATIVE:
            scenario['host'] = 'Подготовка G1–G8, затем «Сценарии подробно», ID '+sid+'. ' + NEGATIVE[sid][0][1]
            scenario['client'] = 'Полный порядок и команды — «Сценарии подробно», ID '+sid+', либо HTML-инструкция. Для S01–S03 также отдельная БЕЗ-ДРУГА.html.'
        else:
            scenario['conditions'], scenario['client'] = SCENARIO_SETUP[sid]
            scenario['host'] = 'Подготовка P1–P7; адрес A1–A9; затем полный порядок «Сценарии подробно», ID '+sid+'. Все настройки: MINECRAFT\\config\\skycraft.properties, G2–G5.'
        if sid == 'S02':
            scenario['limit'] = 'Ожидаемый отказ netcheck и /join на закрытый порт, затем восстановление своего мира. Второй игрок не проверяется.'
        if sid == 'S04':
            scenario['limit'] = 'Требуется среда разработки; в friend-kit NOT_RUN. Другой игрок и совместная игра не проверяются.'
        if sid == 'N02':
            scenario['expected'] = 'Полный цикл по IP проходит. При наличии собственного DNS-имени повторный цикл по имени проходит; без имени DNS-часть BLOCKED.'
    (OUT/'scenarios.json').write_text(json.dumps(scenarios,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    book=Workbook();book.remove(book.active)
    start_rows=[
        ['0','Автоматические проверки','Открыть SkyCraft-Хост.exe, SkyCraft-Клиент.exe или SkyCraft-Без-друга.exe из нового комплекта. Папки определяются автоматически. Полный порядок: docs/YS-AUTOTEST.md. Ниже остаются подробные ручные проверки HUD/блоков/разных условий.'],
        ['1','Существующий SkyCraft работает','Распаковать friend-kit вне игры → Установить обновление.cmd → Проверить сборку.cmd → одиночный HUD/движение.'],
        ['2','Первый тест с другом','В одной LAN — L01. В разных сетях — I01: WAN IP/проброс TCP или доступный IPv6. Адрес получить по вкладке Адрес хоста, не наугад.'],
        ['3','Читать конкретный сценарий','В Сценарии подробно включить фильтр столбца ID, оставить нужный ID. Там полный порядок с командами и нужным Run ID. Подготовка и Адрес хоста объясняют источники переменных.'],
        ['4','Ввод команд','T в Skyrim открывает чат Minecraft; /... вводить там. Команды Get-/Test- вводить в Windows PowerShell. Все подписи ENDPOINT/IP заменить своим значением.'],
        ['5','Логи и результат','Прогоны — наблюдения/метрики, Условия прогонов — адрес/версии/режимы. В конце два ZIP с одним ID. PASS TCP не означает PASS игрового сценария.'],
        ['6','Запуск без сети','Проблема HUD/движения: Собрать отчёт.cmd режим 1. Он соберёт AUTO JSONL; ручной Run ID ещё не нужен.'],
        ['7','Где файл настроек','Сначала вкладка «Файлы и команды», G1–G5. Параметры join/network.* находятся в MINECRAFT\\config\\skycraft.properties. Пустой join = строка join= без текста после знака =.'],
        ['8','Друг сейчас недоступен','Вкладка «Без друга»: B0–B10 по порядку. Готовые команды, адреса loopback, пять отчётов и возврат DIRECT. Отдельная HTML: БЕЗ-ДРУГА.html.'],
    ]
    start=page(book,'Начать',['№','Ситуация','Порядок'],start_rows,[7,36,105],'Start')
    page(book,'Файлы и команды',['Код','Действие','Пошагово','Ожидание'],BASICS,[8,32,98,82],'FilesCommands')
    page(book,'Подготовка',['Код','Действие','Что сделать','Ожидание'],PREPARATION,[8,30,90,65],'Preparation')
    page(book,'Адрес хоста',['Код','Ветка','Точные действия','Что получить / ограничения'],ADDRESSES,[8,32,96,70],'Addresses')
    page(book,'Настройка сети',['Код','Режим','Точные действия','Что проверить'],TRANSPORTS,[8,30,90,68],'Transports')
    page(book,'Без друга',['Код','Тест','Действия на одном ПК','Ожидание и отчёт'],SOLO,[8,36,110,85],'SinglePC')
    page(book,'Шаги',['Шаг','Этап','Хост','Клиент','Что фиксировать'],STEPS,[8,28,70,74,64],'Steps')
    matrix_rows=[[s[k] for k in ('id','name','transport','minutes','conditions','expected')] for s in scenarios]
    matrix=page(book,'Матрица',['ID','Сценарий','Транспорт','Минуты','Условия','PASS'],matrix_rows,[9,38,24,10,68,75],'Scenarios')
    detailed=scenario_rows(scenarios)
    page(book,'Сценарии подробно',['ID','Шаг','Этап','Хост / один ПК','Клиент','Ожидание / запись'],detailed,[8,10,28,75,75,70],'Detailed')
    run_headers=['Сценарий','Run ID','Результат','TCP','Вход, мс','Ping сред., мс','Ping макс., мс','Минуты факт','Обрывы','Движение обоим','Блоки обоим','Leave / rejoin','ZIP хоста','ZIP клиента','Шаг / UTC / сообщение / заметки']
    condition_headers=['Run ID','Начало UTC','Сеть','Адрес:порт','Источник адреса / внутренний и внешний порт','Версии хоста','Режимы хоста / версии / профиль / порядок','Версии клиента','Режимы клиента / версии / профиль / порядок','Транспорт клиента','Прокси клиента:порт (без пароля)']
    run_rows=[];condition_rows=[]
    for s in scenarios:
        for repeat in range(1,4):
            run_id=f'{s["id"]}_{repeat:02}';run_rows.append([s['id'],run_id,'NOT_RUN']+[None]*(len(run_headers)-3));condition_rows.append([run_id]+[None]*(len(condition_headers)-1))
    runs=page(book,'Прогоны',run_headers,run_rows,[10,16,17,17,14,16,16,14,11,18,17,22,32,32,65],'Runs')
    page(book,'Условия прогонов',condition_headers,condition_rows,[16,24,16,32,56,40,62,40,62,23,34],'RunConditions')
    choices={'C':'PASS,FAIL,BLOCKED,PARTIAL,NOT_RUN','D':'OK,REFUSED,TIMEOUT,DNS_ERROR,AUTH_ERROR,NOT_RUN','J':'YES,NO,NOT_RUN','K':'YES,NO,NOT_RUN','L':'OK,FAIL,NOT_RUN'}
    for col,items in choices.items():
        validation=DataValidation(type='list',formula1='"'+items+'"',allow_blank=True);validation.showErrorMessage=True;runs.add_data_validation(validation);validation.add(f'{col}2:{col}1001')
    for col in 'EFGHI':
        validation=DataValidation(type='decimal',operator='greaterThanOrEqual',formula1=0,allow_blank=True);validation.showErrorMessage=True;runs.add_data_validation(validation);validation.add(f'{col}2:{col}1001')
    for result,color in [('PASS','DDF3E4'),('FAIL','FCE4E4'),('BLOCKED','FFF0C2'),('PARTIAL','DFECFC')]:
        runs.conditional_formatting.add('C2:C76',FormulaRule(formula=[f'C2="{result}"'],fill=PatternFill('solid',fgColor=color)))
    page(book,'Откуда данные',['Поле','Источник','Куда записать'],DATA,[28,92,86],'Sources')
    page(book,'Сбор отчёта',['Код','Этап','Каждое поле / действие','Где результат'],REPORTS,[8,35,110,85],'ReportFields')
    page(book,'Разбор сбоя',['Симптом','Точные проверки','Трактовка'],FAILURES,[30,94,84],'Failures')
    for row,s in enumerate(scenarios,2):
        target=next(i+2 for i,r in enumerate(detailed) if r[0]==s['id']);matrix.cell(row,1).hyperlink=f"#'Сценарии подробно'!A{target}";matrix.cell(row,1).font=Font(color='087F8C',underline='single')
    for row,name in enumerate(book.sheetnames[1:],start.max_row+3):
        start.cell(row,2,'Открыть: '+name).hyperlink=f"#'{name}'!A1";start.cell(row,2).font=Font(color='087F8C',underline='single')
    path=OUT/'SkyCraft-Network-Tests.xlsx';book.save(path)
    original_headers=['ID','Когда','Сценарий','Транспорт','Что делает хост','Что делает клиент','Условия','Ожидаемый результат','Минуты','Ограничение']
    csv_file('scenarios.csv',original_headers,[[s[k] for k in ('id','stage','name','transport','host','client','conditions','expected','minutes','limit')] for s in scenarios])
    csv_file('runs-template.csv',run_headers,run_rows);csv_file('conditions-template.csv',condition_headers,condition_rows)
    csv_file('detailed-steps.csv',['ID','Шаг','Этап','Хост / один ПК','Клиент','Ожидание'],detailed)
    csv_file('solo-steps.csv',['Код','Этап','На одном ПК','Ожидание'],SOLO)
    render_guides(scenarios,detailed)
    check=load_workbook(path)
    assert len(check.worksheets)==14 and check['Матрица'].max_row==26 and check['Прогоны'].max_row==76
    assert all(check['Прогоны'].cell(row,3).value=='NOT_RUN' for row in range(2,77))
    assert len(set(r[0] for r in detailed))==25 and all(len(r)==6 for r in detailed)
    assert 'Передать доступный клиенту адрес' not in (ROOT/'docs/YS-TESTING.md').read_text(encoding='utf-8')
    print(f'DETAILED TEST KIT PASS: 25 scenarios, {len(detailed)} detailed rows, 75 empty runs, 14 tabs, full scenario HTML and solo guide')

if __name__=='__main__':main()
